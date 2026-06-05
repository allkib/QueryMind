"""
Import and export helpers for the wells dataset and query results.

Covers three flows:

- ``import_wells_csv`` — validates an uploaded CSV against the expected schema
  (required columns, numeric coercion, row cap) before it replaces the local
  dataset, so a malformed upload can never poison later queries.
- ``result_to_csv_bytes`` — the simple CSV download of a result table.
- ``result_to_workbook_bytes`` — a richer, multi-sheet ``.xlsx`` report bundling
  the question/metadata, the data, the rendered chart image, and the generated
  code, so a non-technical stakeholder gets a self-contained, auditable artifact
  from one click. ``openpyxl`` is imported lazily inside that function to keep
  the common import path light.
"""

from __future__ import annotations

import base64
import io
from datetime import datetime, timezone
from typing import Any, Dict, List

import pandas as pd

from schema import DATA_PATH

EXPECTED_COLUMNS = [
    "well_id",
    "month",
    "production_bbl",
    "field",
    "operator",
    "depth_ft",
]

MAX_IMPORT_ROWS = 100_000


def import_wells_csv(file_storage: Any) -> Dict[str, Any]:
    """
    Validate and replace the local wells dataset from an uploaded CSV file.
    """
    if not file_storage or not file_storage.filename:
        raise ValueError("No file uploaded.")

    df = pd.read_csv(file_storage)
    if df.empty:
        raise ValueError("CSV file is empty.")

    df.columns = [str(col).strip() for col in df.columns]
    missing = [col for col in EXPECTED_COLUMNS if col not in df.columns]
    if missing:
        raise ValueError(
            f"CSV is missing required columns: {', '.join(missing)}. "
            f"Expected: {', '.join(EXPECTED_COLUMNS)}"
        )

    df = df[EXPECTED_COLUMNS].copy()
    df["production_bbl"] = pd.to_numeric(df["production_bbl"], errors="coerce")
    df["depth_ft"] = pd.to_numeric(df["depth_ft"], errors="coerce")

    if df["production_bbl"].isna().any() or df["depth_ft"].isna().any():
        raise ValueError(
            "production_bbl and depth_ft must be numeric in every row."
        )

    if len(df) > MAX_IMPORT_ROWS:
        raise ValueError(f"CSV exceeds maximum of {MAX_IMPORT_ROWS:,} rows.")

    DATA_PATH.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(DATA_PATH, index=False)

    return {
        "rows": len(df),
        "columns": EXPECTED_COLUMNS,
        "path": str(DATA_PATH),
    }


def result_to_csv_bytes(result: Dict[str, Any]) -> bytes:
    """
    Convert a query result payload into CSV bytes for download.
    """
    if not result or not result.get("columns"):
        raise ValueError("No result data to export.")

    columns: List[str] = [str(c) for c in result["columns"]]
    rows = result.get("rows") or []
    df = pd.DataFrame(rows, columns=columns)
    buffer = io.StringIO()
    df.to_csv(buffer, index=False)
    return buffer.getvalue().encode("utf-8")


# ── Rich Excel workbook export ────────────────────────────────────────────
_HEADER_FILL = "0A2540"  # navy
_HEADER_FONT_COLOR = "FFFFFF"
_LABEL_FILL = "E8F2FA"  # light blue
_MAX_COL_WIDTH = 60


def _decode_chart_image(chart_image: str | None) -> bytes | None:
    """Decode a data-URL / base64 PNG string into raw bytes."""
    if not chart_image or not isinstance(chart_image, str):
        return None
    payload = chart_image.strip()
    if payload.startswith("data:"):
        # data:image/png;base64,XXXX
        _, _, payload = payload.partition(",")
    if not payload:
        return None
    try:
        return base64.b64decode(payload)
    except (ValueError, base64.binascii.Error):
        return None


def result_to_workbook_bytes(payload: Dict[str, Any]) -> bytes:
    """
    Build a multi-sheet Excel workbook from a query result payload:

    - Summary: question, metadata, and export details
    - Results: the data table (styled header, auto-sized columns)
    - Chart: the rendered chart image (when provided)
    - Code: the generated analysis code/SQL

    Returns the .xlsx file as bytes.
    """
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    result = payload.get("result") or {}
    if not result or not result.get("columns"):
        raise ValueError("No result data to export.")

    question = str(payload.get("question") or "").strip()
    code = str(payload.get("code") or "").strip()
    chart_image = payload.get("chart_image")
    chart_type = str(payload.get("chart_type") or "").strip()

    columns: List[str] = [str(c) for c in result["columns"]]
    rows = result.get("rows") or []

    header_fill = PatternFill("solid", fgColor=_HEADER_FILL)
    header_font = Font(bold=True, color=_HEADER_FONT_COLOR)
    label_fill = PatternFill("solid", fgColor=_LABEL_FILL)
    label_font = Font(bold=True)
    mono_font = Font(name="Consolas", size=10)
    title_font = Font(bold=True, size=14, color=_HEADER_FILL)

    wb = Workbook()

    # ── Summary sheet ──────────────────────────────────────────────
    summary = wb.active
    summary.title = "Summary"
    summary["A1"] = "QueryMind Analysis Report"
    summary["A1"].font = title_font
    exported_at = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    meta_rows = [
        ("Question", question or "(not provided)"),
        ("Rows returned", len(rows)),
        ("Columns", len(columns)),
        ("Chart type", chart_type or "n/a"),
        ("Exported", exported_at),
    ]
    for offset, (label, value) in enumerate(meta_rows, start=3):
        cell_label = summary.cell(row=offset, column=1, value=label)
        cell_label.fill = label_fill
        cell_label.font = label_font
        cell_label.alignment = Alignment(vertical="top")
        summary.cell(row=offset, column=2, value=value).alignment = Alignment(
            wrap_text=True, vertical="top"
        )
    summary.column_dimensions["A"].width = 18
    summary.column_dimensions["B"].width = 80

    # ── Results sheet ──────────────────────────────────────────────
    results_ws = wb.create_sheet("Results")
    for col_idx, col_name in enumerate(columns, start=1):
        cell = results_ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="left")
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, col_name in enumerate(columns, start=1):
            results_ws.cell(row=row_idx, column=col_idx, value=row.get(col_name))

    # Auto-size columns based on content width.
    for col_idx, col_name in enumerate(columns, start=1):
        max_len = len(str(col_name))
        for row in rows:
            value = row.get(col_name)
            if value is not None:
                max_len = max(max_len, len(str(value)))
        results_ws.column_dimensions[get_column_letter(col_idx)].width = min(
            max_len + 2, _MAX_COL_WIDTH
        )
    results_ws.freeze_panes = "A2"

    # ── Chart sheet ────────────────────────────────────────────────
    chart_ws = wb.create_sheet("Chart")
    chart_ws["A1"] = "Chart"
    chart_ws["A1"].font = title_font
    image_bytes = _decode_chart_image(chart_image)
    if image_bytes:
        try:
            img = XLImage(io.BytesIO(image_bytes))
            chart_ws.add_image(img, "A3")
        except Exception:
            chart_ws["A3"] = "Chart image could not be embedded."
    else:
        chart_ws["A3"] = "No chart was available for this result."

    # ── Code sheet ─────────────────────────────────────────────────
    code_ws = wb.create_sheet("Code")
    code_ws["A1"] = "Generated analysis code"
    code_ws["A1"].font = title_font
    code_lines = code.splitlines() or ["(no code generated)"]
    for offset, line in enumerate(code_lines, start=3):
        cell = code_ws.cell(row=offset, column=1, value=line)
        cell.font = mono_font
    code_ws.column_dimensions["A"].width = 100

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out.getvalue()
